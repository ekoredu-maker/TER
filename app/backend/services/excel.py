from __future__ import annotations

import csv
import io
import re
import uuid
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .timeutils import iso_kst


ALIASES = {
    "orderNo": ["순번"],
    "dept": ["소속", "부서", "기관", "담당부서"],
    "sourceDept": ["신청당시부서"],
    "position": ["직급", "직위"],
    "name": ["성명", "이름", "출장자", "직원명", "신청자"],
    "purpose": ["출장목적", "목적", "업무내용", "출장내용"],
    "period": ["출장기간", "출장일시", "일정", "출장기간시각"],
    "durationText": ["일수기간", "일수/기간", "기간"],
    "startDate": ["시작일", "출발일", "출장시작일", "출장일자", "일자"],
    "startTime": ["시작시간", "출발시간"],
    "endDate": ["종료일", "도착일", "출장종료일"],
    "endTime": ["종료시간", "도착시간"],
    "destination": ["출장지", "장소", "방문지", "도착지"],
    "mealProvided": ["식사제공여부", "식사제공", "중식제공", "제공식"],
    "signer": ["서명또는날인", "서명", "날인"],
    "coTravelers": ["복수출장자", "동행자", "동행출장자", "출장인원"],
    "tripType": ["출장종류", "출장구분", "관내외구분"],
    "vehicleUse": ["공용차량이용여부", "공용차량_x000d_이용여부", "공용차량"],
    "approvalStatus": ["결재상태", "승인상태"],
    "deleteStatus": ["삭제여부"],
}

NO_SETTLEMENT_RE = re.compile(r"여비\s*부지급|정산\s*불요|여비없음")


def _normalize_text(value: Any) -> str:
    return re.sub(r"[\s()~\-_:./]", "", str(value or "")).lower()


def _maybe_date(value: Any) -> str:
    if value in (None, ""):
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    s = str(value).strip()
    m = re.search(r"(20\d{2})[.\-/년 ]\s*(\d{1,2})[.\-/월 ]\s*(\d{1,2})", s)
    if m:
        return f"{m.group(1)}-{int(m.group(2)):02d}-{int(m.group(3)):02d}"
    return ""


def _maybe_time(value: Any) -> str:
    if value in (None, ""):
        return ""
    if isinstance(value, datetime):
        return value.strftime("%H:%M")
    s = str(value).strip()
    m = re.search(r"(\d{1,2})[:시](\d{2})?", s)
    if m:
        return f"{int(m.group(1)):02d}:{int(m.group(2) or 0):02d}"
    return ""


def _extract_period(value: Any) -> dict[str, str]:
    s = re.sub(r"\s+", " ", str(value or "").replace("\n", " ")).strip()
    dates = [
        f"{m.group(1)}-{int(m.group(2)):02d}-{int(m.group(3)):02d}"
        for m in re.finditer(r"(20\d{2})[.\-/년 ]\s*(\d{1,2})[.\-/월 ]\s*(\d{1,2})", s)
    ]
    times = [
        f"{int(m.group(1)):02d}:{int(m.group(2) or 0):02d}"
        for m in re.finditer(r"(\d{1,2})[:시](\d{2})?", s)
    ]
    return {
        "startDate": dates[0] if dates else "",
        "endDate": dates[1] if len(dates) > 1 else (dates[0] if dates else ""),
        "startTime": times[0] if times else "",
        "endTime": times[1] if len(times) > 1 else (times[0] if times else ""),
    }


def _mapping(header: list[Any]) -> dict[str, int]:
    normalized = [_normalize_text(v) for v in header]
    result: dict[str, int] = {}

    # Exact alias match first. This prevents short aliases from stealing a wider header.
    for key, aliases in ALIASES.items():
        alias_norm = [_normalize_text(alias) for alias in aliases]
        for i, cell in enumerate(normalized):
            if cell and cell in alias_norm:
                result[key] = i
                break

    # Fallback to substring matching only for still-unmapped fields.
    for key, aliases in ALIASES.items():
        if key in result:
            continue
        alias_norm = [_normalize_text(alias) for alias in aliases]
        for i, cell in enumerate(normalized):
            if cell and any(alias and alias in cell for alias in alias_norm):
                result[key] = i
                break
    return result


def _detect_header(rows: list[list[Any]]) -> tuple[int, dict[str, int]] | None:
    best: tuple[int, dict[str, int]] | None = None
    best_score = -1
    for idx, row in enumerate(rows[:20]):
        mapping = _mapping(row)
        if len(mapping) > best_score:
            best_score = len(mapping)
            best = (idx, mapping)
    return best if best and best_score >= 4 else None


def _approved(value: Any) -> bool:
    s = str(value or "").strip()
    return not s or bool(re.search(r"완결|승인|결재완료", s))


def _deleted(value: Any) -> bool:
    s = str(value or "").strip()
    if not s:
        return False
    if re.fullmatch(r"미삭제|N|NO|아니오|정상|0|FALSE", s, flags=re.I):
        return False
    if re.fullmatch(r"Y|YES|예|삭제|삭제됨|1|TRUE", s, flags=re.I):
        return True
    return True


def _trip_scope(trip_type: Any) -> str:
    s = re.sub(r"\s+", "", str(trip_type or ""))
    if "관내" in s:
        return "intra"
    if "관외" in s:
        return "inter"
    return "unknown"


def _settlement_policy(purpose: Any, trip_type: Any) -> tuple[bool, str]:
    if NO_SETTLEMENT_RE.search(str(purpose or "")):
        return False, "no_expense"
    scope = _trip_scope(trip_type)
    if scope == "intra":
        return False, "intra_default_exempt"
    if scope == "inter":
        return True, "inter_required"
    return True, "unknown_review"


def _cell(row: list[Any], mapping: dict[str, int], key: str, default: Any = "") -> Any:
    idx = mapping.get(key)
    return row[idx] if idx is not None and idx < len(row) else default


def _row_to_trip(row: list[Any], mapping: dict[str, int], filename: str) -> dict[str, Any]:
    period = _extract_period(_cell(row, mapping, "period"))
    start_date = _maybe_date(_cell(row, mapping, "startDate")) or period["startDate"]
    end_date = _maybe_date(_cell(row, mapping, "endDate")) or period["endDate"] or start_date
    start_time = _maybe_time(_cell(row, mapping, "startTime")) or period["startTime"]
    end_time = _maybe_time(_cell(row, mapping, "endTime")) or period["endTime"]
    purpose = str(_cell(row, mapping, "purpose")).strip()
    trip_type = str(_cell(row, mapping, "tripType")).strip()
    settlement_required, settlement_policy = _settlement_policy(purpose, trip_type)
    now = iso_kst()
    return {
        "id": f"trip_{uuid.uuid4().hex}",
        "orderNo": str(_cell(row, mapping, "orderNo")).strip(),
        "dept": str(_cell(row, mapping, "dept")).strip(),
        "sourceDept": str(_cell(row, mapping, "sourceDept")).strip(),
        "position": str(_cell(row, mapping, "position")).strip(),
        "name": str(_cell(row, mapping, "name")).strip(),
        "coTravelers": str(_cell(row, mapping, "coTravelers")).strip(),
        "startDate": start_date,
        "startTime": start_time,
        "endDate": end_date,
        "endTime": end_time,
        "destination": str(_cell(row, mapping, "destination")).strip(),
        "purpose": purpose,
        "mealProvided": str(_cell(row, mapping, "mealProvided")).strip(),
        "signer": str(_cell(row, mapping, "signer")).strip(),
        "tripType": trip_type,
        "tripScope": _trip_scope(trip_type),
        "settlementPolicy": settlement_policy,
        "vehicleUse": str(_cell(row, mapping, "vehicleUse")).strip(),
        "approvalStatus": str(_cell(row, mapping, "approvalStatus")).strip(),
        "deleteStatus": str(_cell(row, mapping, "deleteStatus")).strip(),
        "durationText": str(_cell(row, mapping, "durationText")).strip(),
        "settlementRequired": settlement_required,
        "source": filename,
        "createdAt": now,
        "updatedAt": now,
    }


def _read_rows(filename: str, data: bytes) -> list[list[Any]]:
    suffix = Path(filename).suffix.lower()
    if suffix == ".csv":
        text = data.decode("utf-8-sig", errors="replace")
        return [list(row) for row in csv.reader(io.StringIO(text))]
    if suffix not in {".xlsx", ".xlsm"}:
        raise ValueError("지원 형식은 .xlsx, .xlsm, .csv 입니다.")
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    return [list(row) for row in ws.iter_rows(values_only=True)]


def parse_trip_file(filename: str, data: bytes) -> dict[str, Any]:
    rows = _read_rows(filename, data)
    detected = _detect_header(rows)
    if not detected:
        return {
            "preview": [],
            "warnings": ['헤더 행을 찾지 못했습니다. 상단 20행 안에 "성명, 출장목적, 출장기간/일시, 출장지"를 확인해 주세요.'],
        }

    header_idx, mapping = detected
    preview: list[dict[str, Any]] = []
    warnings: list[str] = []
    blank_streak = 0

    for row_no, row in enumerate(rows[header_idx + 1 :], start=header_idx + 2):
        if not row or all(str(v or "").strip() == "" for v in row):
            blank_streak += 1
            if blank_streak >= 3:
                break
            continue
        blank_streak = 0
        trip = _row_to_trip(row, mapping, filename)
        if not trip["name"] and not trip["purpose"] and not trip["destination"]:
            continue
        if not trip["name"]:
            warnings.append(f"{row_no}행: 신청자/성명이 비어 있어 제외")
            continue
        if _deleted(trip["deleteStatus"]):
            warnings.append(f'{row_no}행: 삭제여부가 "{trip["deleteStatus"]}"라서 제외')
            continue
        if not _approved(trip["approvalStatus"]):
            warnings.append(f'{row_no}행: 결재상태가 "{trip["approvalStatus"]}"라서 제외')
            continue

        if trip["settlementPolicy"] == "no_expense":
            warnings.append(f"{row_no}행: 여비부지급/정산불요 문구가 있어 정산불요로 표시")
        elif trip["settlementPolicy"] == "intra_default_exempt":
            warnings.append(f"{row_no}행: 관내출장으로 확인되어 기본 정산불요로 표시")
        elif trip["settlementPolicy"] == "unknown_review":
            warnings.append(f"{row_no}행: 관내·관외 구분을 확인하지 못해 정산대상으로 임시 표시")

        preview.append(trip)

    return {"preview": preview, "warnings": warnings}


# Copyright 2026@박주가리교감 All rights reserved.
